# -*- coding: utf-8 -*-
"""Exporta registros (de la app o consolidados) a un .xlsx con el mismo
formato que la bitacora de oficina, con el logo de JMF y encabezado.
La columna A (ID) queda oculta y es la que usa merge_bitacoras.py para
no duplicar filas al consolidar.

NOTA sobre filas: por el logo/titulo, los encabezados de columna ya NO
estan en la fila 1 sino en la fila HEADER_ROW, y los datos empiezan en
DATA_START_ROW. merge_bitacoras.py lee estas mismas constantes, asi que
si cambias esto aqui, no hace falta tocar nada mas alla.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
except ImportError:  # Pillow no disponible: se exporta sin logo, sin tronar
    XLImage = None
    PILImage = None

FONT = "Arial"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE_DIR, "assets", "jmf_logo.jpg")

TITULO_ROW = 1
HEADER_ROW = 3
DATA_START_ROW = 4

HEADERS = [
    ("ID", 0),                      # columna oculta, no borrar
    ("Tecnico", 14),
    ("Fecha/Hora", 18),
    ("Nombre del proyecto", 18),
    ("Sondeo / Pozo", 12),
    ("Este (UTM)", 14),
    ("Norte (UTM)", 14),
    ("Prof. desde (m)", 10),
    ("Prof. hasta (m)", 10),
    ("Nivel de roca (m)", 12),
    ("Nivel de agua (m)", 12),
    ("Simbolo USCS", 12),
    ("Nombre de grupo (ASTM D2487)", 22),
    ("Color", 16),
    ("Humedad", 10),
    ("Consistencia / Compacidad", 16),
    ("Plasticidad", 14),
    ("Resistencia en seco", 14),
    ("Dilatancia", 12),
    ("Tenacidad", 10),
    ("Estructura", 14),
    ("Forma de particulas", 14),
    ("Tamano max. particula", 16),
    ("Reaccion con HCl", 12),
    ("Olor", 14),
    ("Formacion geologica / Nombre local", 18),
    ("Comentarios adicionales", 22),
    ("DESCRIPCION ESTANDAR (generada en campo)", 55),
]

# Orden de campos del dict 'registro' que corresponde a cada columna (desde la 4a, tras ID/Tecnico/Fecha)
_CAMPOS_ORDEN = [
    "proyecto", "sondeo", "coord_este", "coord_norte",
    "prof_desde", "prof_hasta", "nivel_roca", "nivel_agua",
    "simbolo_uscs", "nombre_grupo", "color", "humedad",
    "consistencia_compacidad", "plasticidad", "resistencia_seca",
    "dilatancia", "tenacidad", "estructura", "forma_particulas",
    "tamano_max_particula", "reaccion_hcl", "olor",
    "formacion", "comentarios",
]


def _estilos():
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return header_fill, header_font, border, wrap, center


def exportar_a_xlsx(registros: list, ruta_salida: str):
    """registros: lista de dicts como los que devuelve db.listar_registros()."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitacora"
    header_fill, header_font, border, wrap, center = _estilos()
    n_cols = len(HEADERS)

    # --- Fila de titulo con logo JMF ---
    ws.merge_cells(start_row=TITULO_ROW, start_column=3, end_row=TITULO_ROW, end_column=n_cols)
    titulo = ws.cell(row=TITULO_ROW, column=3,
                      value="JMF INGENIERIA & CONSTRUCCION  —  Bitacora Geotecnica (ASTM D2488)")
    titulo.font = Font(name=FONT, bold=True, size=13, color="1F4E78")
    titulo.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[TITULO_ROW].height = 40

    if XLImage is not None and PILImage is not None and os.path.exists(LOGO_PATH):
        w0, h0 = PILImage.open(LOGO_PATH).size
        img = XLImage(LOGO_PATH)
        img.height = 40
        img.width = int(40 * w0 / h0)
        ws.add_image(img, f"A{TITULO_ROW}")

    for j, (h, w) in enumerate(HEADERS, start=1):
        c = ws.cell(row=HEADER_ROW, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
        if w:
            ws.column_dimensions[get_column_letter(j)].width = w

    ws.freeze_panes = f"B{DATA_START_ROW}"
    ws.row_dimensions[HEADER_ROW].height = 32

    for i, reg in enumerate(registros, start=DATA_START_ROW):
        fila = [reg.get("id", ""), reg.get("tecnico", ""), reg.get("fecha_hora", "")]
        fila += [reg.get(campo, "") for campo in _CAMPOS_ORDEN]
        fila.append(reg.get("descripcion_generada", ""))
        for j, val in enumerate(fila, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = border
            cell.alignment = wrap

    # columna ID oculta (se conserva para poder deduplicar, no para leerla a simple vista)
    ws.column_dimensions["A"].hidden = True
    ws.sheet_view.showGridLines = False

    wb.save(ruta_salida)
    return ruta_salida
